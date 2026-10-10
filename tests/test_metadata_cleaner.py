import base64
import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image, PngImagePlugin
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pypdf import PdfWriter, PdfReader
from pypdf.generic import DictionaryObject, NameObject, ArrayObject, DecodedStreamObject
from mutagen.id3 import ID3, TIT2, TPE1
from mutagen.mp3 import MP3
from fastapi.testclient import TestClient

from metadata_cleaner import clean_file, CleaningRejected
import upload_handler
import web_app


def make_image(path, orientation=1):
    exif = Image.Exif()
    exif[315] = 'Private Author'
    exif[272] = 'Camera Model'
    exif[274] = orientation
    kwargs = {'exif': exif}
    if path.suffix == '.png':
        info = PngImagePlugin.PngInfo()
        info.add_text('Comment', 'Private comment')
        kwargs['pnginfo'] = info
    Image.new('RGB', (24, 16), 'red').save(path, **kwargs)


@pytest.mark.parametrize('suffix', ['.jpg', '.png'])
def test_image_metadata_removed_without_pixel_changes(tmp_path, suffix):
    source, output = tmp_path / ('original' + suffix), tmp_path / ('clean' + suffix)
    make_image(source)
    original = source.read_bytes()
    report = clean_file(source, output)
    assert source.read_bytes() == original
    assert {'Artist', 'Model'} <= set(report['removed_fields'])
    assert 'Artist' not in report['remaining_metadata']
    with Image.open(source) as before, Image.open(output) as after:
        assert before.tobytes() == after.tobytes()
        assert not after.getexif()
        if suffix == '.png':
            assert not after.text
    assert 'not a guarantee' in report['notice']


@pytest.mark.parametrize('suffix', ['.jpg', '.png'])
def test_orientation_refused_and_failed_output_removed(tmp_path, suffix):
    source, output = tmp_path / ('original' + suffix), tmp_path / ('clean' + suffix)
    make_image(source, 6)
    original = source.read_bytes()
    with pytest.raises(CleaningRejected, match='orientation'):
        clean_file(source, output)
    assert not output.exists()
    assert source.read_bytes() == original


@pytest.mark.parametrize('suffix, factory', [('.docx', Document), ('.xlsx', Workbook), ('.pptx', Presentation)])
def test_office_only_core_changes_and_remains_readable(tmp_path, suffix, factory):
    source, output = tmp_path / ('original' + suffix), tmp_path / ('clean' + suffix)
    document = factory()
    if suffix == '.xlsx':
        document.properties.creator = 'Private Author'
        document.active['A1'] = '=1+2'
    else:
        document.core_properties.author = 'Private Author'
        if suffix == '.docx':
            document.add_paragraph('Keep this content')
        else:
            document.slides.add_slide(document.slide_layouts[6])
    document.save(source)
    original = source.read_bytes()
    report = clean_file(source, output)
    assert 'Author' in report['removed_fields']
    assert source.read_bytes() == original
    with ZipFile(source) as before, ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != 'docProps/core.xml':
                assert before.read(name) == after.read(name)
    if suffix == '.xlsx':
        with_output = load_workbook(output)
        assert with_output.active['A1'].value == '=1+2'
        with_output.close()
    elif suffix == '.docx':
        assert Document(output).paragraphs[0].text == 'Keep this content'
    else:
        assert len(Presentation(output).slides) == 1


@pytest.mark.parametrize('suffix, factory', [('.docx', Document), ('.xlsx', Workbook), ('.pptx', Presentation)])
def test_signed_office_rejected(tmp_path, suffix, factory):
    source, output = tmp_path / ('signed' + suffix), tmp_path / ('clean' + suffix)
    factory().save(source)
    with ZipFile(source, 'a') as archive:
        archive.writestr('_xmlsignatures/sig1.xml', '<signature/>')
    original = source.read_bytes()
    with pytest.raises(CleaningRejected, match='Signed'):
        clean_file(source, output)
    assert not output.exists()
    assert source.read_bytes() == original


def make_pdf(path, feature=None):
    writer = PdfWriter()
    page = writer.add_blank_page(width=100, height=100)
    stream = DecodedStreamObject()
    stream.set_data(b'0 0 10 10 re S\n')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.add_metadata({'/Author': 'Private Author'})
    if feature:
        writer.root_object[NameObject(feature)] = DictionaryObject()
    with path.open('wb') as destination:
        writer.write(destination)


def test_simple_pdf_cleaned_and_page_content_preserved(tmp_path):
    source, output = tmp_path / 'original.pdf', tmp_path / 'clean.pdf'
    make_pdf(source)
    original = source.read_bytes()
    report = clean_file(source, output)
    assert source.read_bytes() == original
    assert 'Author' in report['removed_fields']
    with source.open('rb') as a, output.open('rb') as b:
        before, after = PdfReader(a), PdfReader(b)
        assert before.pages[0].get_contents().get_data() == after.pages[0].get_contents().get_data()
        assert before.pages[0].mediabox == after.pages[0].mediabox
        assert not after.metadata


@pytest.mark.parametrize('feature', ['/AcroForm', '/Perms', '/Names', '/OpenAction', '/StructTreeRoot'])
def test_complex_pdf_refused(tmp_path, feature):
    source, output = tmp_path / 'original.pdf', tmp_path / 'clean.pdf'
    make_pdf(source, feature)
    with pytest.raises(CleaningRejected):
        clean_file(source, output)
    assert not output.exists()


def test_mp3_removes_tags_preserves_audio(tmp_path):
    source, output = tmp_path / 'original.mp3', tmp_path / 'clean.mp3'
    frames = (b'\xff\xfb\x90\x00' + b'\0' * 413) * 10
    source.write_bytes(frames)
    tags = ID3()
    tags.add(TIT2(encoding=3, text=['Private title']))
    tags.add(TPE1(encoding=3, text=['Private artist']))
    tags.save(source, v1=2)
    original = source.read_bytes()
    report = clean_file(source, output)
    assert source.read_bytes() == original
    assert output.read_bytes() == frames
    assert {'Title', 'Artist'} <= set(report['removed_fields'])
    assert MP3(output).info.sample_rate == MP3(source).info.sample_rate
    assert MP3(output).info.bitrate == MP3(source).info.bitrate


@pytest.mark.parametrize('suffix', ['.mp4', '.mov', '.gif'])
def test_unsupported_cleaning_is_rejected(tmp_path, suffix):
    source, output = tmp_path / ('original' + suffix), tmp_path / ('clean' + suffix)
    source.write_bytes(b'original')
    with pytest.raises(CleaningRejected, match='not supported'):
        clean_file(source, output)
    assert not output.exists()
    assert source.read_bytes() == b'original'


def test_original_and_existing_destination_protected(tmp_path):
    source = tmp_path / 'original.png'
    make_image(source)
    original = source.read_bytes()
    with pytest.raises(CleaningRejected):
        clean_file(source, source)
    output = tmp_path / 'existing.png'
    output.write_bytes(b'do not overwrite')
    with pytest.raises(CleaningRejected):
        clean_file(source, output)
    assert source.read_bytes() == original
    assert output.read_bytes() == b'do not overwrite'


def test_clean_api_response_and_temporaries_removed(tmp_path, monkeypatch):
    source = tmp_path / 'original.png'
    make_image(source)
    original = source.read_bytes()
    temporary_root = tmp_path / 'temporary'
    temporary_root.mkdir()
    original_directory = upload_handler.TemporaryDirectory
    monkeypatch.setattr(upload_handler, 'TemporaryDirectory', lambda **kw: original_directory(dir=temporary_root, **kw))
    with TestClient(web_app.app) as client:
        response = client.post('/api/clean', files={'file': ('original.png', original)})
    assert response.status_code == 200
    assert response.headers['cache-control'] == 'no-store'
    assert list(temporary_root.iterdir()) == []
    assert source.read_bytes() == original
    result = response.json()
    with Image.open(BytesIO(base64.b64decode(result['data']))) as image:
        assert not image.getexif()
        assert image.size == (24, 16)
    assert 'Artist' in result['report']['removed_fields']


def test_clean_api_rejection_and_cleanup(tmp_path, monkeypatch):
    source = tmp_path / 'oriented.jpg'
    make_image(source, 6)
    root = tmp_path / 'temporary'
    root.mkdir()
    original_directory = upload_handler.TemporaryDirectory
    monkeypatch.setattr(upload_handler, 'TemporaryDirectory', lambda **kw: original_directory(dir=root, **kw))
    with TestClient(web_app.app) as client:
        response = client.post('/api/clean', files={'file': ('image.jpg', source.read_bytes())})
        assert response.status_code == 422
        assert 'orientation' in response.json()['detail']
        assert 'data' not in response.json()
        assert client.post('/api/clean', content=b'x', headers={'content-length': str(web_app.MAX_REQUEST_BYTES + 1)}).status_code == 413
    assert list(root.iterdir()) == []


def test_jpeg_auxiliary_provenance_rejected(tmp_path):
    source, output = tmp_path / 'original.jpg', tmp_path / 'clean.jpg'
    make_image(source)
    data = source.read_bytes()
    payload = b'provenance-signature'
    source.write_bytes(data[:2] + b'\xff\xeb' + (len(payload)+2).to_bytes(2, 'big') + payload + data[2:])
    with pytest.raises(CleaningRejected, match='provenance'):
        clean_file(source, output)
    assert not output.exists()


def test_png_unknown_chunk_rejected(tmp_path):
    import zlib
    source, output = tmp_path / 'original.png', tmp_path / 'clean.png'
    make_image(source)
    data = source.read_bytes()
    kind, payload = b'caBX', b'credentials'
    chunk = len(payload).to_bytes(4, 'big') + kind + payload + (zlib.crc32(kind+payload) & 0xffffffff).to_bytes(4, 'big')
    source.write_bytes(data[:33] + chunk + data[33:])
    with pytest.raises(CleaningRejected, match='provenance'):
        clean_file(source, output)
    assert not output.exists()


def test_pdf_document_xmp_removed(tmp_path):
    source, output = tmp_path / 'original.pdf', tmp_path / 'clean.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    xmp = DecodedStreamObject()
    xmp.set_data(b'<x:xmpmeta xmlns:x="adobe:ns:meta/"><private>Private author</private></x:xmpmeta>')
    xmp[NameObject('/Type')] = NameObject('/Metadata')
    xmp[NameObject('/Subtype')] = NameObject('/XML')
    writer.root_object[NameObject('/Metadata')] = writer._add_object(xmp)
    with source.open('wb') as stream:
        writer.write(stream)
    clean_file(source, output)
    assert PdfReader(output).xmp_metadata is None
    assert b'Private author' not in output.read_bytes()


def test_encrypted_pdf_refused(tmp_path):
    source, output = tmp_path / 'encrypted.pdf', tmp_path / 'clean.pdf'
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt('secret')
    with source.open('wb') as stream:
        writer.write(stream)
    with pytest.raises(upload_handler.UploadValidationError):
        clean_file(source, output)
    assert not output.exists()


def test_xlsx_reader_defaults_are_not_reported_as_stored(tmp_path):
    source, output = tmp_path / 'original.xlsx', tmp_path / 'clean.xlsx'
    workbook = Workbook()
    workbook.properties.creator = 'Alice'
    workbook.save(source)
    report = clean_file(source, output)
    assert not {'Author', 'CreationDate', 'ModDate'} & set(report['remaining_metadata'])
    with ZipFile(output) as archive:
        assert b'Alice' not in archive.read('docProps/core.xml')


def test_verification_failure_removes_copy(tmp_path, monkeypatch):
    import metadata_cleaner
    source, output = tmp_path / 'original.png', tmp_path / 'clean.png'
    make_image(source)
    original = source.read_bytes()
    def fail(*args):
        raise CleaningRejected('The image preservation check failed.')
    monkeypatch.setattr(metadata_cleaner, '_verify_pixels', fail)
    with pytest.raises(CleaningRejected):
        clean_file(source, output)
    assert not output.exists()
    assert source.read_bytes() == original


def test_unexpected_clean_api_error_is_private_and_cleans(tmp_path, monkeypatch):
    source = tmp_path / 'original.png'
    make_image(source)
    root = tmp_path / 'temporary'
    root.mkdir()
    original_directory = upload_handler.TemporaryDirectory
    monkeypatch.setattr(upload_handler, 'TemporaryDirectory', lambda **kw: original_directory(dir=root, **kw))
    def fail(*args):
        raise RuntimeError('/secret/internal/path')
    monkeypatch.setattr(web_app, 'clean_file', fail)
    with TestClient(web_app.app) as client:
        response = client.post('/api/clean', files={'file': ('image.png', source.read_bytes())})
    assert response.status_code == 500
    assert '/secret' not in response.text
    assert list(root.iterdir()) == []


def test_exif_colour_information_rejected(tmp_path):
    source, output = tmp_path / 'colour.jpg', tmp_path / 'clean.jpg'
    exif = Image.Exif()
    exif[40961] = 65535
    Image.new('RGB', (10, 10)).save(source, exif=exif)
    with pytest.raises(CleaningRejected, match='colour'):
        clean_file(source, output)
    assert not output.exists()


@pytest.mark.parametrize('suffix, brand', [('.mp4', b'isom'), ('.mov', b'qt  ')])
def test_video_cleaning_disabled_at_api(suffix, brand):
    import struct
    def box(kind, payload=b''):
        return struct.pack('>I4s', len(payload)+8, kind) + payload
    data = box(b'ftyp', brand + b'\0' * 4) + box(b'moov')
    with TestClient(web_app.app) as client:
        response = client.post('/api/clean', files={'file': ('video' + suffix, data)})
    assert response.status_code == 422
    assert 'disabled' in response.json()['detail']
    assert 'data' not in response.json()
