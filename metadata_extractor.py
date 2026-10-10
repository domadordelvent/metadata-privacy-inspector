import json
import subprocess


from pathlib import Path
from PIL import Image, ExifTags
from gps_extractor import extract_gps

from pypdf import PdfReader

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from mutagen.mp3 import MP3


def extract_image_metadata(file_path: Path) -> dict:
    metadata = {}

    with Image.open(file_path) as image:
        metadata["format"] = image.format
        metadata["width"] = image.width
        metadata["height"] = image.height

        exif = image.getexif()

        gps = extract_gps(exif)

        if gps:
            metadata["GPS"] = gps

        for tag_id, value in exif.items():
            tag = ExifTags.TAGS.get(tag_id, str(tag_id))
            metadata[tag] = value

    return metadata



def extract_pdf_metadata(file_path: Path) -> dict:
    metadata = {}

    reader = PdfReader(file_path)
    pdf_metadata = reader.metadata

    if not pdf_metadata:
        return metadata

    for key, value in pdf_metadata.items():
        metadata[key.lstrip("/")] = value

    return metadata

def extract_docx_metadata(file_path: Path) -> dict:
    metadata = {}

    document = Document(file_path)
    properties = document.core_properties

    fields = {
        "Title": properties.title,
        "Subject": properties.subject,
        "Author": properties.author,
        "Keywords": properties.keywords,
        "Comments": properties.comments,
        "LastModifiedBy": properties.last_modified_by,
        "CreationDate": properties.created,
        "ModDate": properties.modified,
        "Category": properties.category,
    }

    for key, value in fields.items():
        if value not in (None, ""):
            metadata[key] = value
    return metadata


def extract_xlsx_metadata(file_path: Path) -> dict:
    metadata = {}

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        properties = workbook.properties

        fields = {
            "Title": properties.title,
            "Subject": properties.subject,
            "Author": properties.creator,
            "Keywords": properties.keywords,
            "Comments": properties.description,
            "LastModifiedBy": properties.lastModifiedBy,
            "CreationDate": properties.created,
            "ModDate": properties.modified,
            "Category": properties.category,
        }

        for key, value in fields.items():
            if value not in (None, ""):
                metadata[key] = value

        return metadata
    finally:
        workbook.close()


def extract_pptx_metadata(file_path: Path) -> dict:
    metadata = {}

    presentation = Presentation(file_path)
    properties = presentation.core_properties

    fields = {
        "Title": properties.title,
        "Subject": properties.subject,
        "Author": properties.author,
        "Keywords": properties.keywords,
        "Comments": properties.comments,
        "LastModifiedBy": properties.last_modified_by,
        "CreationDate": properties.created,
        "ModDate": properties.modified,
        "Category": properties.category,
    }

    for key, value in fields.items():
        if value not in (None, ""):
            metadata[key] = value

    return metadata

def extract_mp3_metadata(file_path: Path) -> dict:
    metadata = {}

    audio = MP3(file_path)

    tag_mapping = {
        "TIT2": "Title",
        "TPE1": "Artist",
        "TALB": "Album",
        "TPE2": "AlbumArtist",
        "TCOM": "Composer",
        "TCON": "Genre",
        "TDRC": "Date",
        "TSSE": "Encoder",
    }

    if audio.tags:
        for key, value in audio.tags.items():
            field = tag_mapping.get(key, key)

            if key.startswith("TXXX:"):
                field = key.split(":", 1)[1].capitalize()

            metadata[field] = str(value)

    if audio.info:
        metadata["Duration"] = audio.info.length
        metadata["Bitrate"] = audio.info.bitrate
        metadata["SampleRate"] = audio.info.sample_rate

    return metadata

def extract_video_metadata(file_path: Path) -> dict:
    metadata = {}

    result = subprocess.run(
        [
            "ffprobe",
            "-v", "quiet",
            "-print_format", "json",
            "-show_format",
            "-show_streams",
            str(file_path),
        ],
        capture_output=True,
        text=True,
        check=True,
        timeout=15,
    )

    data = json.loads(result.stdout)

    format_info = data.get("format", {})

    metadata["Duration"] = format_info.get("duration")
    metadata["Bitrate"] = format_info.get("bit_rate")

    tags = format_info.get("tags", {})

    tag_mapping = {
        "title": "Title",
        "artist": "Artist",
        "author": "Author",
        "comment": "Comment",
        "creation_time": "CreationDate",
        "encoder": "Encoder",
        "location": "Location",
        "location-eng": "Location",
    }

    for key, value in tags.items():
        field = tag_mapping.get(key.lower(), key)
        metadata[field] = value

    return {
        key: value
        for key, value in metadata.items()
        if value not in (None, "")
    }